"""
Universal multi-format document parser supporting PDF, Microsoft Word (.docx),
Excel spreadsheets (.xlsx, .csv), Markdown, Code files, and plain text.
Extracts clean, normalized text while preserving page/sheet numbers and source metadata.
"""

import io
import re
import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Union
import pypdf

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


@dataclass
class DocumentPage:
    """Represents a single page, sheet, or segment of an ingested document."""
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    """Represents a fully parsed document with metadata and page-level text."""
    filename: str
    doc_id: str
    file_type: str
    total_pages: int
    pages: List[DocumentPage] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        """Returns the concatenated text across all pages."""
        return "\n\n".join(page.text for page in self.pages)


class DocumentParser:
    """Parser that ingests files from filesystem paths or in-memory file buffers across all formats."""

    @staticmethod
    def clean_text(text: str) -> str:
        """Normalizes unicode whitespace, cleans artifacts, and strips excessive newlines."""
        if not text:
            return ""
        # Replace non-breaking spaces and tabs
        text = text.replace("\xa0", " ").replace("\t", " ")
        # Replace 3 or more newlines with 2
        text = re.sub(r"\n{3,}", "\n\n", text)
        # Replace multiple horizontal spaces with single space
        text = re.sub(r"[ ]{2,}", " ", text)
        return text.strip()

    def parse_pdf(self, file_source: Union[str, Path, io.BytesIO], filename: str) -> ParsedDocument:
        """Parses a PDF document page by page."""
        reader = pypdf.PdfReader(file_source)
        pages: List[DocumentPage] = []

        for idx, page in enumerate(reader.pages):
            raw_text = page.extract_text() or ""
            cleaned = self.clean_text(raw_text)
            if cleaned:
                pages.append(DocumentPage(page_number=idx + 1, text=cleaned))

        if not pages:
            pages.append(DocumentPage(page_number=1, text="[Document contained no extractable text]"))

        return ParsedDocument(
            filename=filename,
            doc_id=Path(filename).stem,
            file_type="pdf",
            total_pages=len(reader.pages),
            pages=pages,
        )

    def parse_docx(self, file_source: Union[str, Path, io.BytesIO], filename: str) -> ParsedDocument:
        """Parses a Microsoft Word (.docx) document including paragraphs and tables."""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required to parse .docx files")

        if isinstance(file_source, (str, Path)):
            doc = docx.Document(str(file_source))
        else:
            doc = docx.Document(file_source)

        sections: List[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                sections.append(text)

        # Extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(row_cells):
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                sections.append("\n".join(table_rows))

        combined_text = "\n\n".join(sections)
        cleaned = self.clean_text(combined_text)

        # Estimate page count (roughly 400 words / 2400 chars per page)
        est_pages = max(1, len(cleaned) // 2400 + (1 if len(cleaned) % 2400 else 0))

        # Split into estimated pages if lengthy
        pages: List[DocumentPage] = []
        if est_pages > 1:
            chunk_len = len(cleaned) // est_pages
            for i in range(est_pages):
                start = i * chunk_len
                end = (i + 1) * chunk_len if i < est_pages - 1 else len(cleaned)
                pages.append(DocumentPage(page_number=i + 1, text=cleaned[start:end].strip()))
        else:
            pages.append(DocumentPage(page_number=1, text=cleaned or "[Empty Word Document]"))

        return ParsedDocument(
            filename=filename,
            doc_id=Path(filename).stem,
            file_type="docx",
            total_pages=len(pages),
            pages=pages,
        )

    def parse_excel(self, file_source: Union[str, Path, io.BytesIO], filename: str) -> ParsedDocument:
        """Parses an Excel spreadsheet (.xlsx) sheet by sheet."""
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl is required to parse .xlsx files")

        wb = openpyxl.load_workbook(file_source, data_only=True)
        pages: List[DocumentPage] = []

        for idx, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            sheet_rows = []
            for row in sheet.iter_rows(values_only=True):
                non_empty = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if non_empty:
                    sheet_rows.append(" | ".join(non_empty))

            if sheet_rows:
                sheet_text = f"--- Sheet: {sheet_name} ---\n" + "\n".join(sheet_rows)
                pages.append(DocumentPage(page_number=idx, text=self.clean_text(sheet_text)))

        if not pages:
            pages.append(DocumentPage(page_number=1, text="[Empty Excel Spreadsheet]"))

        return ParsedDocument(
            filename=filename,
            doc_id=Path(filename).stem,
            file_type="xlsx",
            total_pages=len(pages),
            pages=pages,
        )

    def parse_csv(self, file_source: Union[str, Path, io.BytesIO], filename: str) -> ParsedDocument:
        """Parses a CSV file and formats tabular data clearly."""
        if isinstance(file_source, io.BytesIO):
            content_bytes = file_source.read()
            try:
                text_content = content_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text_content = content_bytes.decode("latin-1", errors="replace")
        else:
            path = Path(file_source)
            try:
                text_content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                text_content = path.read_text(encoding="latin-1", errors="replace")

        reader = csv.reader(io.StringIO(text_content))
        rows = [" | ".join(r) for r in reader if any(cell.strip() for cell in r)]
        cleaned = self.clean_text("\n".join(rows))

        return ParsedDocument(
            filename=filename,
            doc_id=Path(filename).stem,
            file_type="csv",
            total_pages=1,
            pages=[DocumentPage(page_number=1, text=cleaned or "[Empty CSV]")],
        )

    def parse_text(self, file_source: Union[str, Path, io.BytesIO], filename: str) -> ParsedDocument:
        """Parses plain text, Markdown, or Code source files."""
        if isinstance(file_source, io.BytesIO):
            content_bytes = file_source.read()
            try:
                raw_text = content_bytes.decode("utf-8")
            except UnicodeDecodeError:
                raw_text = content_bytes.decode("latin-1", errors="replace")
        else:
            path = Path(file_source)
            try:
                raw_text = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                raw_text = path.read_text(encoding="latin-1", errors="replace")

        cleaned = self.clean_text(raw_text)
        ext = Path(filename).suffix.lower().lstrip(".")

        return ParsedDocument(
            filename=filename,
            doc_id=Path(filename).stem,
            file_type=ext if ext else "txt",
            total_pages=1,
            pages=[DocumentPage(page_number=1, text=cleaned or "[Empty Document]")],
        )

    def parse(self, file_source: Union[str, Path, io.BytesIO], filename: Optional[str] = None) -> ParsedDocument:
        """Determines document format by extension and routes to appropriate parser."""
        if filename is None:
            if isinstance(file_source, (str, Path)):
                filename = Path(file_source).name
            else:
                filename = "uploaded_document"

        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            return self.parse_pdf(file_source, filename)
        elif ext in [".docx", ".doc"]:
            return self.parse_docx(file_source, filename)
        elif ext in [".xlsx", ".xls"]:
            return self.parse_excel(file_source, filename)
        elif ext == ".csv":
            return self.parse_csv(file_source, filename)
        else:
            # Code, Markdown, TXT, JSON, RTF, etc.
            return self.parse_text(file_source, filename)
